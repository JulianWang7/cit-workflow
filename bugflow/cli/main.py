"""bugfix-batch — 批量 Bug 修复编排器（单会话编排 + 子Agent 派发）。

读取 Excel bug 列表 → 逐个驱动单会话编排工作流 → 回写结果 → 故障隔离。

Usage:
    bugfix-batch --excel bugs.xlsx --server 252 --code-root /home/android/source
    bugfix-batch --bug-id 12345 --server 252 --code-root /home/android/source
    bugfix-batch --excel bugs.xlsx --dry-run  # 只打印计划不执行
    bugfix-batch --excel bugs.xlsx --legacy-3session  # 旧版 3 会话拆分模式

工作流（每个 bug，单会话编排模式 — 默认）:
    单会话 orchestrator (主会话持有跨阶段状态):
      1. set_workspace
      2. Agent(bugfix-analyst) → ROOT_CAUSE_FOUND     (独立 context，只读分析)
      3. Agent(bugfix-worker)  → PATCH_DONE            (独立 context，读写修复)
      4. Agent(bugfix-reviewer) → REVIEW_PASS/FAIL     (独立 context，只读审查)
      5. compile → verify → submit (主会话直接执行)    (需要持有 git/设备状态)
      6. SUBMIT_DONE

工作流（每个 bug，旧版 3 会话模式 --legacy-3session）:
    fixer-pre:  reproduce → analyze → modify   (PATCH_DONE)
    reviewer:   review                        (REVIEW_PASS/FAIL)
    fixer-post: compile → verify → submit      (SUBMIT_DONE)

故障隔离: 任一会话失败/超时 → 记录错误 → git reset → 跳到下一个 bug
"""

from __future__ import annotations

import argparse
import os
import re
import sys
import time
from dataclasses import dataclass, field
from pathlib import Path

from bugflow.credentials import load_zcode_credentials
from bugflow.cli.adapters.base import AgentAdapter, AgentResult, SessionConfig
from bugflow.cli.adapters.zcode_adapter import ZCodeAdapter
from bugflow.core.config import load_server_config, list_servers
from bugflow.core.git import git_exec, reset_to_clean, get_status
from bugflow.core.skill_feedback import infer_subsystem
from bugflow.core import batch_state, fix_log


# ── Bug 数据结构 ──────────────────────────────────────────────
@dataclass
class BugItem:
    """Excel 中一行 bug。"""

    bug_id: str
    title: str
    description: str
    reproduce_steps: str
    code_root_hint: str = ""     # 可能涉及的源码路径提示
    status: str = "pending"      # pending / running / fixed / failed / skipped
    session_log: list[dict] = field(default_factory=list)  # 每步的 {step, success, marker, elapsed, error}
    error: str = ""
    fix_files: list[str] = field(default_factory=list)  # 修改的源码文件路径
    root_cause_category: str = ""  # 根因子所属子系统（如 audio_subsystem）

    @property
    def is_done(self) -> bool:
        return self.status in ("fixed", "failed", "skipped")


# ── Excel 读写 ────────────────────────────────────────────────
def load_bugs_from_excel(excel_path: str) -> list[BugItem]:
    """从 Excel 读取 bug 列表。

    必须列: Bug ID / Title / Description
    可选列: Reproduce Steps / Code Root Hint / Status
    （真实场景通常无 Reproduce Steps，由 AI 根据标题+描述推导复现命令）
    """
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("❌ 需要 openpyxl: pip install openpyxl", file=sys.stderr)
        sys.exit(1)

    wb = load_workbook(excel_path, read_only=True)
    ws = wb.active

    # 读表头，建立列名 → 列号映射
    headers = {}
    for col_idx, cell in enumerate(next(ws.iter_rows(min_row=1, max_row=1))):
        if cell.value:
            headers[str(cell.value).strip().lower()] = col_idx

    # 列名别名
    col_map = {}
    for key, aliases in [
        ("bug_id", ["bug id", "id", "编号", "bug_id"]),
        ("title", ["title", "标题", "bug标题"]),
        ("description", ["description", "描述", "bug描述", "详情"]),
        ("reproduce", ["reproduce", "reproduce steps", "复现步骤", "复现"]),
        ("code_hint", ["code root hint", "code_hint", "源码提示", "路径提示"]),
        ("status", ["status", "状态"]),
    ]:
        for alias in aliases:
            if alias in headers:
                col_map[key] = headers[alias]
                break

    # 必须列检查
    required = ["bug_id", "title", "description"]
    missing = [r for r in required if r not in col_map]
    if missing:
        print(f"❌ Excel 缺少必须列: {missing}。已有列: {list(headers.keys())}", file=sys.stderr)
        sys.exit(1)

    bugs = []
    for row in ws.iter_rows(min_row=2):
        def get(key: str) -> str:
            if key not in col_map:
                return ""
            val = row[col_map[key]].value
            return str(val).strip() if val else ""

        bug_id = get("bug_id")
        if not bug_id:
            continue

        bugs.append(BugItem(
            bug_id=bug_id,
            title=get("title"),
            description=get("description"),
            reproduce_steps=get("reproduce"),
            code_root_hint=get("code_hint"),
            status=get("status") or "pending",
        ))

    wb.close()
    return bugs


def write_results_to_excel(excel_path: str, bugs: list[BugItem]) -> None:
    """回写结果到 Excel（中文状态 + 颜色 + error 列）。

    状态映射（对齐 batch_state + SKILL.md 枚举）:
      fixed   → 已修复已验证（绿色）
      failed  → 可复现未修复（黄色）
      skipped → 可复现未修复（黄色）
      pending → 待处理（无色）
      running → 进行中（无色）
    """
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill

    # 状态映射 + 颜色 (ARGB hex without alpha)
    STATUS_MAP = {
        "fixed":   ("已修复已验证",  "C6EFCE"),  # 绿色
        "failed":  ("可复现未修复", "FFEB9C"),  # 黄色
        "skipped": ("可复现未修复", "FFEB9C"),  # 黄色
        "pending": ("待处理",       ""),         # 无色
        "running": ("进行中",       ""),         # 无色
    }

    wb = load_workbook(excel_path)
    ws = wb.active

    # 读表头，1-indexed 列号
    headers = {}
    for cell in next(ws.iter_rows(min_row=1, max_row=1)):
        if cell.value:
            headers[str(cell.value).strip().lower()] = cell.column  # 1-indexed

    # status 列：已有则用，没有则新建
    status_col = headers.get("status")
    if status_col is None:
        status_col = ws.max_column + 1
        ws.cell(row=1, column=status_col, value="修复状态")

    # error 列
    error_col = headers.get("error")
    if error_col is None:
        error_col = ws.max_column + 1
        ws.cell(row=1, column=error_col, value="备注")

    # bug_id → 列号（注意 0 是 falsy，不能用 or 链）
    id_col = None
    for alias in ("bug id", "id", "编号", "bug_id"):
        if alias in headers:
            id_col = headers[alias]
            break
    if id_col is None:
        print("⚠️ 找不到 Bug ID 列，无法回写", file=sys.stderr)
        return

    # bug_id → row 映射
    id_to_row = {}
    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        val = row[id_col - 1].value  # row 是 0-indexed list, id_col 是 1-indexed
        if val:
            id_to_row[str(val).strip()] = row_idx

    for bug in bugs:
        row_idx = id_to_row.get(bug.bug_id)
        if row_idx:
            cn_status, color = STATUS_MAP.get(bug.status, (bug.status, ""))
            cell = ws.cell(row=row_idx, column=status_col, value=cn_status)
            if color:
                cell.fill = PatternFill(
                    start_color=color, end_color=color, fill_type="solid",
                )
            ws.cell(row=row_idx, column=error_col, value=bug.error[:500])

    wb.save(excel_path)
    wb.close()


# ── 会话 goal 模板 ────────────────────────────────────────────
def _build_orchestrator_prompt(
    bug: BugItem, server: str, code_root: str, device_serial: str,
) -> str:
    """单会话编排提示词 — 主会话用 Agent 工具派发 3 个子 Agent，然后直接编译验证提交。

    主会话 = orchestrator（持有 bug 信息、跨阶段决策、git/设备状态）
    子 Agent = bugfix-analyst / bugfix-worker / bugfix-reviewer（独立 context，专用工具白名单）
    """
    # 复现步骤
    if bug.reproduce_steps:
        repro_section = f"- 复现步骤: {bug.reproduce_steps}\n"
        repro_hint = ""
    else:
        repro_section = (
            "- 复现步骤: 未提供。请根据 bug 标题和描述推导能在设备上复现的 adb shell 命令序列\n"
        )
        repro_hint = (
            "   常见模式：app crash → am start；SELinux denied → 触发对应操作；"
            "系统属性 → setprop + am start\n"
        )

    return (
        f"你是 Android Bug 修复编排者。请用子 Agent 协作修复 bug {bug.bug_id}。\n\n"
        f"## Bug 信息\n"
        f"- ID: {bug.bug_id}\n"
        f"- 标题: {bug.title}\n"
        f"- 描述: {bug.description}\n"
        f"{repro_section}"
        f"- 源码提示: {bug.code_root_hint or '无，需搜索定位'}\n\n"
        f"## 环境\n"
        f"- SSH 服务器: {server}\n"
        f"- 代码根目录: {code_root}\n"
        f"- 设备 serial: {device_serial or '自动选择'}\n\n"
        f"## 编排流程（严格按序执行）\n\n"
        f"### 步骤 1: 设置工作区\n"
        f"调用 set_workspace(server={server}, code_root={code_root})\n\n"
        f"### 步骤 2: 派发分析子 Agent\n"
        f"用 Agent 工具派发 bugfix-analyst:\n"
        f"  subagent_type: \"bugfix-analyst\"\n"
        f"  prompt: \"分析 bug {bug.bug_id}。Bug 标题: {bug.title}。描述: {bug.description}。\n"
        f"          先 set_workspace(server={server}, code_root={code_root})，\n"
        f"          然后搜索源码定位根因，输出根因报告。\n"
        f"          搜索规则：search_code_tool 内部即 rg，禁用 grep -r 搜源码。\"\n"
        f"等待子 Agent 返回。检查返回中是否包含 ROOT_CAUSE_FOUND。\n"
        f"如果未包含 ROOT_CAUSE_FOUND，输出 ANALYZE_FAIL 并停止。\n\n"
        f"### 步骤 3: 派发修复子 Agent\n"
        f"用 Agent 工具派发 bugfix-worker:\n"
        f"  subagent_type: \"bugfix-worker\"\n"
        f"  prompt: \"修复 bug {bug.bug_id}。根因分析: <粘贴 analyst 返回的根因报告>。\n"
        f"          先 set_workspace(server={server}, code_root={code_root})，\n"
        f"          然后最小改动修复源码，输出 PATCH_DONE。\"\n"
        f"等待返回。检查 PATCH_DONE。\n\n"
        f"### 步骤 4: 派发审查子 Agent\n"
        f"用 Agent 工具派发 bugfix-reviewer:\n"
        f"  subagent_type: \"bugfix-reviewer\"\n"
        f"  prompt: \"审查 bug {bug.bug_id} 的 patch。根因报告: <粘贴 analyst 返回的根因报告>。\n"
        f"          先 set_workspace(server={server}, code_root={code_root})，\n"
        f"          然后 git_op(action='diff') 读取改动。\n"
        f"          审查 patch 是否针对根因报告中的源码位置（不是在别处掩盖症状）。\n"
        f"          输出 REVIEW_PASS 或 REVIEW_FAIL。\"\n"
        f"等待返回。如果 REVIEW_FAIL，回到步骤 3 重新修复（最多重试 2 次）。\n\n"
        f"### 步骤 5: 编译验证提交（主会话直接执行，不派发子 Agent）\n"
        f"5a. 编译:\n"
        f"  - 先 run_command('ls out/build-*.ninja 2>/dev/null | head -1') 确定 variant\n"
        f"  - compile_module(module=<从 diff 推导模块路径>, background=True)\n"
        f"  - compile_status(task_id=...) 轮询（首次等 ≥30 秒，后续每 30 秒）\n"
        f"  - 编译成功 → 继续；失败 → 输出 COMPILE_FAIL + 原因\n\n"
        f"5b. 验证:\n"
        f"  - verify(verify_commands=<复现命令>, logcat_filter=<关键词>)\n"
        f"  - symptom_gone=true 且 patch 修改了根因报告中的源码位置 → VERIFY_PASS → 继续\n"
        f"  - symptom_gone=true 但 patch 未修改根因位置 → VERIFY_WARN → 不提交，回 analyze 重新分析\n"
        f"  - symptom_gone=false → VERIFY_FAIL + 原因\n"
        f"  - 判断 root_cause_addressed: git_op(action='status') 列改动文件，与根因报告'源码位置'比对\n"
        f"  - 休眠唤醒型 bug 验证时 adb 可能断连，用 serial_read 兜底抓日志\n"
        f"  - ⚠️ verify_commands 含 echo mem/power_test 时，必须先执行休眠测试安全协议\n"
        f"    （关 autosleep: echo off > /sys/power/autosleep; 设 RTC alarm: echo +15 > /sys/class/rtc/rtc0/wakealarm;\n"
        f"     禁 unbind USB 控制器; 禁 power_test 0 + echo mem 组合），详见 reproduce skill\n"
        f"  {repro_hint}\n"
        f"5c. 提交:\n"
        f"  - git_op(action='commit', commit_message='{bug.bug_id} [Description] {bug.title} [Solution] <修改简述>')\n"
        f"  - git_op(action='branch') 获取当前分支名\n"
        f"  - run_command('git remote -v') 查看可用远端\n"
        f"  - git_op(action='push', push_target='<remote> HEAD:refs/for/<branch>')\n"
        f"  - push 失败 → run_command('git format-patch -1 HEAD -o /tmp/patches/') 保存补丁\n\n"
        f"### 步骤 6: 输出最终标记\n"
        f"输出 SUBMIT_DONE（或 COMPILE_FAIL / VERIFY_FAIL / VERIFY_WARN + 原因）\n\n"
        f"## 输出要求\n"
        f"每步完成时输出对应标记。标记纯文本，单独一行，不加 markdown 格式符。\n"
        f"子 Agent 返回的分析报告和审查报告，请完整保留在你的输出中（用于后续校验）。"
    )


def _build_fixer_pre_prompt(bug: BugItem, server: str, code_root: str, device_serial: str) -> str:
    """fixer-pre 会话: reproduce → analyze → modify。"""
    # 复现步骤可能未提供（真实场景），AI 需根据标题+描述推导
    if bug.reproduce_steps:
        repro_line = f"- 复现步骤: {bug.reproduce_steps}\n"
        repro_task = (
            f"2. 在设备上复现 bug，抓 logcat 证据（使用 reproduce 工具，reproduce_commands 用上面的复现步骤）。\n"
            f"   休眠唤醒型 bug 可用 adb_suspend 触发，adb 断连时 serial_read 兜底抓日志。\n"
        )
    else:
        repro_line = "- 复现步骤: 未提供。请根据 bug 标题和描述，推导能在设备上复现该 bug 的 adb shell 命令序列\n"
        repro_task = (
            f"2. 根据 bug 标题和描述推导复现命令，在设备上复现 bug，抓 logcat 证据（使用 reproduce 工具）\n"
            f"   常见模式：app crash → am start；SELinux denied → 触发对应操作；"
            "系统属性 → setprop + am start；"
            "休眠唤醒型 → adb_suspend + serial 兜底（adb 断连时 serial_read 抓日志）；"
            "UI 操作型 → ui_snapshot/ui_tap_text/ui_wait_text（不用裸 input tap）\n"
        )
    return (
        f"你是 Android Bug 修复 agent。请按以下步骤修复 bug {bug.bug_id}。\n\n"
        f"## Bug 信息\n"
        f"- ID: {bug.bug_id}\n"
        f"- 标题: {bug.title}\n"
        f"- 描述: {bug.description}\n"
        f"{repro_line}"
        f"- 源码提示: {bug.code_root_hint or '无，需搜索定位'}\n\n"
        f"## 环境\n"
        f"- SSH 服务器: {server}\n"
        f"- 代码根目录: {code_root}\n"
        f"- 设备 serial: {device_serial or '自动选择'}\n\n"
        f"## 任务\n"
        f"1. 先调用 set_workspace 设置工作区（server={server}, code_root={code_root}）\n"
        f"{repro_task}"
        f"3. 分析根因，搜索源码定位问题（使用 search_code_tool / locate_files_tool / read_file）\n"
        f"   搜索规则：search_code_tool 内部即 rg（ripgrep），禁用 grep -r 搜源码；"
        "子系统 skill 中若有 grep -r 示例一律以 search_code_tool 为准\n"
        f"4. 最小改动修复（使用 edit_file / write_file）\n"
        f"5. 确认改动（git_op diff）\n\n"
        f"## 输出要求\n"
        f"每步完成时输出对应标记。最终必须在最后一行单独输出 PATCH_DONE（或 REPRODUCE_FAIL + 原因）。\n"
        f"标记格式：纯文本，不要加 markdown 格式符（不要 **、`、#），单独一行。\n"
        f"不要 commit，只改工作区文件。"
    )


def _build_reviewer_prompt(bug: BugItem, server: str, code_root: str,
                           root_cause_report: str = "") -> str:
    """reviewer 会话: review patch。传入根因报告供交叉验证。"""
    report_section = ""
    if root_cause_report:
        report_section = (
            f"\n## 根因分析报告（来自 analyze 阶段）\n"
            f"{root_cause_report}\n"
        )
    return (
        f"你是 Android Bug 修复 reviewer。请审查 bug {bug.bug_id} 的 patch。\n\n"
        f"## Bug 信息\n"
        f"- ID: {bug.bug_id}\n"
        f"- 标题: {bug.title}\n"
        f"- 描述: {bug.description}\n"
        f"{report_section}"
        f"\n## 环境\n"
        f"- SSH 服务器: {server}\n"
        f"- 代码根目录: {code_root}\n\n"
        f"## 任务\n"
        f"1. 先调用 set_workspace 设置工作区\n"
        f"2. 调用 git_op(action='diff') 读取 fixer-pre 产出的 patch\n"
        f"3. 调用 git_op(action='status') 确认改动文件列表\n"
        f"4. 审查 patch 的正确性、完整性、副作用\n"
        f"5. **根因交叉验证（必做）**: 将 patch 改动文件与根因报告中的源码位置/因果链比对，\n"
        f"   确认 patch 针对的是根因(L2/L3)而非在别处掩盖症状(L1)\n"
        f"6. 输出 REVIEW_PASS 或 REVIEW_FAIL（附原因）\n\n"
        f"## 输出要求\n"
        f"最终必须在最后一行单独输出 REVIEW_PASS 或 REVIEW_FAIL。\n"
        f"标记格式：纯文本，不要加 markdown 格式符（不要 **、`、#），单独一行。\n\n"
        f"## 注意\n"
        f"- 只读不写，不要 edit_file / write_file\n"
        f"- 重点检查：改动是否真正修复根因，是否过宽，是否有遗漏\n"
        f"- 如果 patch 文件 ≠ 根因报告源码位置 → REVIEW_FAIL（疑似症状掩盖）"
    )


def _build_fixer_post_prompt(bug: BugItem, server: str, code_root: str, device_serial: str) -> str:
    """fixer-post 会话: compile → verify → submit。"""
    # 复现步骤可能未提供，verify 时需根据描述推导验证命令
    if bug.reproduce_steps:
        repro_line = f"- 复现步骤: {bug.reproduce_steps}\n"
        verify_task = (
            f"4. 推补丁到设备验证（使用 verify 工具，verify_commands 用上面的复现步骤）\n"
            f"   休眠唤醒型 bug 验证时 adb 可能断连，用 serial_read 兜底抓日志\n"
            f"   ⚠️ verify_commands 含 echo mem/power_test 时，必须先执行休眠测试安全协议\n"
            f"   （关 autosleep + 设 RTC alarm + 禁 unbind USB），详见 reproduce skill\n"
        )
    else:
        repro_line = "- 复现步骤: 未提供。请根据 bug 标题和描述推导验证命令\n"
        verify_task = (
            f"4. 推补丁到设备验证（使用 verify 工具，verify_commands 根据 bug 描述推导，"
            f"应与 reproduce 阶段一致）\n"
            f"   休眠唤醒型 bug 验证时 adb 可能断连，用 serial_read 兜底抓日志\n"
            f"   ⚠️ verify_commands 含 echo mem/power_test 时，必须先执行休眠测试安全协议\n"
            f"   （关 autosleep + 设 RTC alarm + 禁 unbind USB），详见 reproduce skill\n"
        )
    return (
        f"你是 Android Bug 修复 agent。请完成 bug {bug.bug_id} 的编译验证提交推送。\n\n"
        f"## Bug 信息\n"
        f"- ID: {bug.bug_id}\n"
        f"- 标题: {bug.title}\n"
        f"- 描述: {bug.description}\n"
        f"{repro_line}\n"
        f"## 环境\n"
        f"- SSH 服务器: {server}\n"
        f"- 代码根目录: {code_root}\n"
        f"- 设备 serial: {device_serial or '自动选择'}\n\n"
        f"## 任务\n"
        f"1. 先调用 set_workspace 设置工作区\n"
        f"2. 调用 git_op(action='diff') 确认 patch 仍在工作区\n"
        f"3. 编译修改的模块（使用 compile_module）\n"
        f"{verify_task}"
        f"5. 验证通过后提交（git_op action='commit'）\n"
        f"6. 推送到 Gerrit（git_op action='push'）：\n"
        f"   - 先 git_op(action='branch') 获取当前分支名\n"
        f"   - 再 run_command('git remote -v') 查看可用远端\n"
        f"   - 构造 push_target: '<remote> HEAD:refs/for/<branch>'\n"
        f"   - git_op(action='push', push_target='<remote> HEAD:refs/for/<branch>')\n\n"
        f"## 输出要求\n"
        f"每步完成时输出对应标记。最终输出 SUBMIT_DONE（或 COMPILE_FAIL / VERIFY_FAIL / VERIFY_WARN + 原因）。\n"
        f"VERIFY_WARN = 症状消失但 patch 未修改根因位置（疑似症状掩盖），不可提交，需回 analyze。\n"
        f"push 失败时不输出 SUBMIT_DONE，输出错误信息。\n"
        f"标记格式：纯文本，不要加 markdown 格式符（不要 **、`、#），单独一行。"
    )


# ── 编排器 ────────────────────────────────────────────────────
# ── 各会话默认超时（秒） ─────────────────────────────────────────
# orchestrator: 单会话含 3 子Agent 派发 + 编译验证提交（子Agent 各自独立 context）
# fixer-pre: reproduce(设备操作+logcat) + analyze(多轮搜索+读文件) + modify(编辑)
# reviewer:  读 diff + 读文件 + 给 verdict（较轻）
# fixer-post: compile(mmm 增量编译) + verify(推设备+重启+重跑) + submit(git commit)
_SESSION_TIMEOUTS = {
    "orchestrator": 3600,  # 60 分钟（单会话含 3 子Agent + 编译验证提交）
    "fixer-pre": 1200,   # 20 分钟
    "reviewer": 600,     # 10 分钟
    "fixer-post": 1800,  # 30 分钟（编译 + 重启 + 验证）
}


class BugfixOrchestrator:
    """批量 Bug 修复编排器。"""

    def __init__(
        self,
        adapter: AgentAdapter,
        server: str,
        code_root: str,
        device_serial: str = "",
        timeout_per_session: int = 0,
        dry_run: bool = False,
        legacy_3session: bool = False,
        resume: bool = False,
    ):
        self.adapter = adapter
        self.server = server
        self.code_root = code_root
        self.device_serial = device_serial
        # 0 = 按会话类型自动（_SESSION_TIMEOUTS）；>0 = 全部用此值
        self.timeout_override = timeout_per_session
        self.dry_run = dry_run
        self.legacy_3session = legacy_3session
        self.resume = resume

    def _timeout_for(self, step_name: str) -> int:
        """按会话类型取超时，或用 --timeout 覆盖。"""
        if self.timeout_override > 0:
            return self.timeout_override
        return _SESSION_TIMEOUTS.get(step_name, 600)

    def _common_env(self) -> dict[str, str]:
        """所有会话共享的环境变量（MCP server 也读这些）。"""
        env = {
            "BUGFIX_SERVER": self.server,
            "BUGFIX_CODE_ROOT": self.code_root,
        }
        if self.device_serial:
            env["BUGFIX_DEVICE_SERIAL"] = self.device_serial
        return env

    def _run_session(self, step_name: str, config: SessionConfig) -> AgentResult:
        """执行一个会话并打印进度。"""
        print(f"\n{'='*60}")
        print(f"  ▶ {step_name}")
        print(f"{'='*60}")
        print(f"  mode: {config.mode}, timeout: {config.timeout_s}s")

        if self.dry_run:
            print(f"  [DRY RUN] goal: {config.goal[:100]}...")
            return AgentResult(success=True, output="[dry run]", marker="DRY_RUN", elapsed_s=0.0)

        result = self.adapter.run_session(config)

        # 打印结果
        status_icon = "✅" if result.success else "❌"
        print(f"  {status_icon} {step_name} 完成: marker={result.marker}, elapsed={result.elapsed_s:.1f}s")
        if result.error:
            print(f"  ⚠️ error: {result.error}")
        # 打印输出最后 5 行
        if result.output:
            lines = result.output.strip().split("\n")
            for line in lines[-5:]:
                print(f"    │ {line}")

        return result

    def _try_update_zentao(self, bug_id: str) -> None:
        """best-effort 回写禅道 bug 状态为 resolved。失败不影响修复结果。"""
        try:
            bid = int(bug_id)
        except (ValueError, TypeError):
            return  # bug_id 不是数字（可能是 Excel 自定义 ID），跳过

        try:
            from bugflow.core import zentao
            comment = "已通过 android-bugfix-flow 自动修复（reproduce→analyze→modify→review→compile→verify→submit）"
            result = zentao.update_bug(bid, status="resolved", comment=comment)
            print(f"  📝 禅道回写: {result}")
        except Exception as e:
            print(f"  ⚠️ 禅道回写失败（不影响修复结果）: {e}", file=sys.stderr)

    def fix_one_bug(self, bug: BugItem) -> None:
        """修复单个 bug: 单会话编排模式（默认）或 3 会话拆分模式（--legacy-3session）。"""
        print(f"\n{'#'*60}")
        print(f"  # Bug {bug.bug_id}: {bug.title}")
        print(f"{'#'*60}")

        # fix_log: 记录修复开始时间（供 log_fix_outcome 算耗时）
        fix_log.start_fix_tracking(str(bug.bug_id))

        try:
            if self.legacy_3session:
                self._fix_one_bug_legacy(bug)
            else:
                self._fix_one_bug_orchestrator(bug)
        finally:
            # fix_log: 记录修复结果
            self._log_fix_outcome(bug)
            fix_log.clear_fix_tracking()

    def _capture_fix_files(self, bug: BugItem) -> None:
        """从 git 获取修改文件列表（best-effort，不阻断主流程）。

        在 session 返回后、reset 前调用。先试已提交的 diff（成功路径），
        为空则试未提交的 status（失败路径，reset 前）。
        """
        if self.dry_run:
            return
        try:
            # 成功路径：agent 已 commit，取最后一次提交的文件
            out, _, ec = git_exec(
                self.server, self.code_root, "diff --name-only HEAD~1 HEAD",
            )
            if ec == 0 and out.strip():
                bug.fix_files = [f.strip() for f in out.splitlines() if f.strip()]
                return
            # 失败路径：未提交的改动（reset 前仍在工作区）
            status_out = get_status(self.server, self.code_root)
            if status_out.strip():
                bug.fix_files = [
                    line.split(maxsplit=1)[-1]
                    for line in status_out.splitlines()
                    if line.strip()
                ]
        except Exception as e:
            print(f"  ⚠️ fix_files 采集失败（不影响主流程）: {e}")

    def _log_fix_outcome(self, bug: BugItem) -> None:
        """记录修复结果到 fix_log.jsonl（best-effort，不阻断主流程）。"""
        try:
            error_step = ""
            if bug.status == "failed" and bug.error:
                # 从 error 字段推断卡在哪一步
                err_lower = bug.error.lower()
                if "compile" in err_lower:
                    error_step = "compile"
                elif "verify" in err_lower:
                    error_step = "verify"
                elif "reproduce" in err_lower:
                    error_step = "reproduce"
                elif "analyze" in err_lower:
                    error_step = "analyze"
                else:
                    error_step = "orchestrator"

            # 从 session_log 提取 gate violations 作为 notes
            notes_parts: list[str] = []
            for entry in bug.session_log:
                if entry.get("step") == "gate_validation" and entry.get("violations"):
                    notes_parts.append(f"门控违规: {'; '.join(entry['violations'])}")

            # 推断根因子子系统（若未被预先设置）
            if not bug.root_cause_category:
                bug.root_cause_category = infer_subsystem(bug.title, bug.fix_files)

            fix_log.log_fix_outcome(
                bug_id=str(bug.bug_id),
                root_cause_category=bug.root_cause_category,
                fix_files=bug.fix_files,
                symptom_gone=(bug.status == "fixed") if bug.status in ("fixed", "failed") else None,
                error_step=error_step,
                notes=" | ".join(notes_parts) if notes_parts else "",
            )
        except Exception as e:
            print(f"  ⚠️ fix_log 记录失败（不影响主流程）: {e}")

    def _fix_one_bug_orchestrator(self, bug: BugItem) -> None:
        """单会话编排模式: 主会话派发 3 子Agent + 直接编译验证提交。"""
        bug.status = "running"
        env = self._common_env()

        # ── git reset 确保干净起点 ──
        if not self.dry_run:
            try:
                reset_to_clean(self.server, self.code_root)
                print(f"  ✅ git reset → 干净工作区")
            except Exception as e:
                print(f"  ⚠️ git reset 失败（可能工作区已干净）: {e}")

        # ── 单会话编排 ──
        config = SessionConfig(
            goal=_build_orchestrator_prompt(
                bug, self.server, self.code_root, self.device_serial,
            ),
            mode="target",
            cwd=self.code_root if Path(self.code_root).is_dir() else "",
            permission_mode="yolo",
            timeout_s=self._timeout_for("orchestrator"),
            env=env,
        )
        result = self._run_session("orchestrator", config)
        bug.session_log.append({
            "step": "orchestrator", "success": result.success,
            "marker": result.marker, "elapsed_s": result.elapsed_s,
            "error": result.error,
        })

        # 采集修改文件（reset 前窗口）
        self._capture_fix_files(bug)

        # 轻量门控校验（不阻断，记录违规）
        violations = self._validate_orchestrator_output(result, bug)
        if violations:
            print(f"  ⚠️ 门控校验发现 {len(violations)} 项违规:")
            for v in violations:
                print(f"     • {v}")

        # 检查 SUBMIT_DONE
        if result.marker == "SUBMIT_DONE":
            bug.status = "fixed"
            print(f"\n  🎉 Bug {bug.bug_id} 修复完成!")
            if not self.dry_run:
                self._try_update_zentao(bug.bug_id)
        else:
            bug.status = "failed"
            bug.error = f"orchestrator 未产出 SUBMIT_DONE (marker={result.marker})"
            if result.error:
                bug.error += f": {result.error}"
            print(f"\n  ❌ Bug {bug.bug_id} 在 orchestrator 失败: {bug.error}")
            # 失败 reset
            if not self.dry_run:
                try:
                    reset_to_clean(self.server, self.code_root)
                except Exception:
                    pass

    def _validate_orchestrator_output(
        self, result: AgentResult, bug: BugItem,
    ) -> list[str]:
        """轻量校验编排器输出，模拟结论门控的 post-hoc 检查。

        返回违规列表（不阻断，记录到 session_log 供 fix_log 统计）。
        检查项:
        1. 根因结论缺少 path:line 证据引用
        2. confidence=high 但含 hedge words（可能/疑似/大概/或许/推测）
        3. category=unknown 但给出了根因猜测
        """
        violations: list[str] = []
        output = result.output
        if not output:
            return violations

        output_lower = output.lower()

        # 检查 1: 如果输出包含根因结论但无 path:line 引用
        # path:line 格式: filename.ext:123（至少一个路径:行号引用）
        has_root_cause = "root_cause" in output_lower or "根因" in output
        has_path_line = bool(re.search(r'[\w/]+\.\w+:\d+', output))
        if has_root_cause and not has_path_line:
            violations.append("根因结论缺少 path:line 证据引用")

        # 检查 2: 如果 confidence=high 但含 hedge words
        if "confidence" in output_lower and "high" in output_lower:
            hedge_words = ["可能", "疑似", "大概", "或许", "推测"]
            found_hedges = [w for w in hedge_words if w in output]
            if found_hedges:
                violations.append(
                    f"confidence=high 但含 hedging 用语: {', '.join(found_hedges)}"
                )

        # 检查 3: category=unknown 但给了根因
        if "category" in output_lower and "unknown" in output_lower:
            if has_root_cause:
                violations.append("category=unknown 不应给出根因猜测")

        # 记录到 session_log
        if violations:
            bug.session_log.append({
                "step": "gate_validation", "violations": violations,
            })

        return violations

    def _fix_one_bug_legacy(self, bug: BugItem) -> None:
        """旧版 3 会话拆分模式: fixer-pre → reviewer → fixer-post。"""
        bug.status = "running"
        env = self._common_env()

        # ── git reset 确保干净起点 ──
        if not self.dry_run:
            try:
                reset_to_clean(self.server, self.code_root)
                print(f"  ✅ git reset → 干净工作区")
            except Exception as e:
                print(f"  ⚠️ git reset 失败（可能工作区已干净）: {e}")

        # ── Session 1: fixer-pre (reproduce → analyze → modify) ──
        pre_config = SessionConfig(
            goal=_build_fixer_pre_prompt(bug, self.server, self.code_root, self.device_serial),
            mode="target",
            cwd=self.code_root if Path(self.code_root).is_dir() else "",
            permission_mode="yolo",
            timeout_s=self._timeout_for("fixer-pre"),
            env=env,
        )
        pre_result = self._run_session("fixer-pre", pre_config)
        bug.session_log.append({
            "step": "fixer-pre", "success": pre_result.success,
            "marker": pre_result.marker, "elapsed_s": pre_result.elapsed_s,
            "error": pre_result.error,
        })

        # 检查 PATCH_DONE
        if pre_result.marker != "PATCH_DONE":
            bug.status = "failed"
            bug.error = f"fixer-pre 未产出 PATCH_DONE (marker={pre_result.marker})"
            if pre_result.error:
                bug.error += f": {pre_result.error}"
            print(f"\n  ❌ Bug {bug.bug_id} 在 fixer-pre 失败: {bug.error}")
            return

        # ── Session 2: reviewer ──
        reviewer_config = SessionConfig(
            goal=_build_reviewer_prompt(bug, self.server, self.code_root,
                                       root_cause_report=pre_result.output),
            mode="target",
            permission_mode="yolo",
            timeout_s=self._timeout_for("reviewer"),
            env=env,
        )
        review_result = self._run_session("reviewer", reviewer_config)
        bug.session_log.append({
            "step": "reviewer", "success": review_result.success,
            "marker": review_result.marker, "elapsed_s": review_result.elapsed_s,
            "error": review_result.error,
        })

        # 检查 REVIEW_PASS
        if review_result.marker != "REVIEW_PASS":
            bug.status = "failed"
            bug.error = f"reviewer 未通过 (marker={review_result.marker})"
            print(f"\n  ❌ Bug {bug.bug_id} review 未通过: {bug.error}")
            # review 失败也 reset，保持工作区干净给下一个 bug
            if not self.dry_run:
                try:
                    reset_to_clean(self.server, self.code_root)
                except Exception:
                    pass
            return

        # ── Session 3: fixer-post (compile → verify → submit) ──
        post_config = SessionConfig(
            goal=_build_fixer_post_prompt(bug, self.server, self.code_root, self.device_serial),
            mode="target",
            permission_mode="yolo",
            timeout_s=self._timeout_for("fixer-post"),
            env=env,
        )
        post_result = self._run_session("fixer-post", post_config)
        bug.session_log.append({
            "step": "fixer-post", "success": post_result.success,
            "marker": post_result.marker, "elapsed_s": post_result.elapsed_s,
            "error": post_result.error,
        })

        # 采集修改文件（reset 前窗口）
        self._capture_fix_files(bug)

        # 检查 SUBMIT_DONE
        if post_result.marker == "SUBMIT_DONE":
            bug.status = "fixed"
            print(f"\n  🎉 Bug {bug.bug_id} 修复完成!")
            # best-effort 回写禅道（bug_id 是数字时才尝试）
            if not self.dry_run:
                self._try_update_zentao(bug.bug_id)
        else:
            bug.status = "failed"
            bug.error = f"fixer-post 未产出 SUBMIT_DONE (marker={post_result.marker})"
            if post_result.error:
                bug.error += f": {post_result.error}"
            print(f"\n  ❌ Bug {bug.bug_id} 在 fixer-post 失败: {bug.error}")
            # 失败 reset
            if not self.dry_run:
                try:
                    reset_to_clean(self.server, self.code_root)
                except Exception:
                    pass

    def run_batch(self, bugs: list[BugItem]) -> list[BugItem]:
        """批量执行所有 pending bug。

        集成 batch_state 断点续跑:
        - --resume: 加载当日 batch_state，跳过已达终态的 bug
        - 每个 bug 修复后更新 batch_state 并落盘
        """
        import datetime

        today = datetime.date.today().isoformat()

        # Phase 0: 加载 batch_state（始终加载，确保 state 有 date 字段可落盘）
        # --resume: 额外过滤掉已达终态的 bug
        state = batch_state.load_batch_state(today)
        if self.resume:
            all_ids = [str(b.bug_id) for b in bugs if b.status == "pending"]
            pending_ids = set(batch_state.get_pending_bugs(state, all_ids))
            skipped_count = len(all_ids) - len(pending_ids)
            if skipped_count > 0:
                print(f"  📋 resume: 跳过 {skipped_count} 个已达终态的 bug")
            # 过滤掉终态 bug
            pending = [b for b in bugs if b.status == "pending" and str(b.bug_id) in pending_ids]
        else:
            pending = [b for b in bugs if b.status == "pending"]

        print(f"\n批量修复: {len(pending)} 个 pending bug (共 {len(bugs)} 个)")

        for i, bug in enumerate(pending, 1):
            print(f"\n[{i}/{len(pending)}] 开始处理 Bug {bug.bug_id}")
            try:
                self.fix_one_bug(bug)
            except KeyboardInterrupt:
                print(f"\n⚠️ 用户中断，Bug {bug.bug_id} 标记为 skipped")
                bug.status = "skipped"
                bug.error = "user interrupted"
                # 更新 batch_state
                self._update_batch_state(state, today, bug)
                break
            except Exception as e:
                # 故障隔离: 异常不传播到下一个 bug
                print(f"\n❌ Bug {bug.bug_id} 异常: {e}")
                bug.status = "failed"
                bug.error = f"exception: {e}"
                # 尝试 reset
                if not self.dry_run:
                    try:
                        reset_to_clean(self.server, self.code_root)
                    except Exception:
                        pass
            finally:
                # 更新 batch_state（每个 bug 处理后，无论成功/失败/中断）
                self._update_batch_state(state, today, bug)

        # 汇总
        fixed = sum(1 for b in bugs if b.status == "fixed")
        failed = sum(1 for b in bugs if b.status == "failed")
        skipped = sum(1 for b in bugs if b.status == "skipped")
        print(f"\n{'='*60}")
        print(f"  批量完成: {len(bugs)} 个 bug")
        print(f"  ✅ fixed: {fixed}")
        print(f"  ❌ failed: {failed}")
        print(f"  ⏭️ skipped: {skipped}")
        print(f"{'='*60}")

        # 打印 batch_state 摘要
        if state:
            try:
                summary = batch_state.state_summary(state)
                print(f"  📊 batch_state: {summary.get('done_count', 0)} 终态 / "
                      f"{summary.get('pending_count', 0)} 待处理 / "
                      f"{summary.get('total', 0)} 总计")
            except Exception:
                pass

        return bugs

    def _update_batch_state(
        self, state: dict, date: str, bug: BugItem,
    ) -> None:
        """更新单 bug 的 batch_state 并落盘（best-effort，不阻断主流程）。

        状态映射:
        - fixed → 已修复已验证（终态，resume 跳过）
        - failed → 可复现未修复（非终态，resume 可重试）
        - skipped(中断) → 可复现未修复（非终态，resume 可重试）

        异常/中断通常是临时问题（SSH 断连/超时/用户中断），不应标记为终态，
        让 --resume 能重试这些 bug。
        """
        try:
            # 映射 BugItem.status → batch_state.fix_status
            status_map = {
                "fixed": "已修复已验证",
                "failed": "可复现未修复",
                "skipped": "可复现未修复",
            }
            fix_status = status_map.get(bug.status, "")

            batch_state.update_bug_state(
                state, str(bug.bug_id),
                title=bug.title,
                module=bug.root_cause_category,
                fix_status=fix_status,
                phase="completed" if bug.status == "fixed" else "attempted",
                notes=bug.error or ("fixed" if bug.status == "fixed" else ""),
            )
            batch_state.save_batch_state(state)
        except Exception as e:
            print(f"  ⚠️ batch_state 更新失败（不影响主流程）: {e}")


# ── CLI 入口 ──────────────────────────────────────────────────
def main():
    parser = argparse.ArgumentParser(
        prog="bugfix-batch",
        description="Android Bug 批量修复编排器",
    )
    parser.add_argument("--excel", help="Bug 列表 Excel 文件路径（与 --bug-id 互斥）")
    parser.add_argument("--bug-id", help="禅道 Bug ID，逗号分隔多个（如 12345,12346）。从禅道自动拉取标题/描述")
    parser.add_argument("--server", required=True, help="SSH 服务器名 (servers.yaml 中定义)")
    parser.add_argument("--code-root", required=True, help="Android 源码根目录")
    parser.add_argument("--device-serial", default="", help="ADB 设备 serial (不传自动选)")
    parser.add_argument("--timeout", type=int, default=0, help="覆盖所有会话超时秒数 (默认 0=按会话类型: orchestrator=3600s/fixer-pre=1200s/reviewer=600s/fixer-post=1800s)")
    parser.add_argument("--dry-run", action="store_true", help="只打印计划不执行")
    parser.add_argument("--legacy-3session", action="store_true", help="使用旧版 3 会话拆分模式（不使用子Agent编排）")
    parser.add_argument("--resume", action="store_true", help="从 batch_state 断点续跑（跳过已达终态的 bug）")
    parser.add_argument("--provider", default="", help="Provider ID 或 model 名子串 (如 'volcengine' / 'DeepSeek')，不传自动选第一个可用的")
    args = parser.parse_args()

    # 验证 --excel / --bug-id 互斥（必须有一个）
    if not args.excel and not args.bug_id:
        print("❌ 必须指定 --excel 或 --bug-id 之一", file=sys.stderr)
        sys.exit(1)
    if args.excel and args.bug_id:
        print("❌ --excel 和 --bug-id 不能同时使用", file=sys.stderr)
        sys.exit(1)

    # 验证 server 配置
    try:
        load_server_config(args.server)
    except Exception as e:
        print(f"❌ 服务器 '{args.server}' 配置不可用: {e}", file=sys.stderr)
        print(f"   可用服务器: {list_servers()}", file=sys.stderr)
        sys.exit(1)

    # 加载 bug 列表
    if args.bug_id:
        from bugflow.core import zentao
        try:
            bug_ids = [int(x.strip()) for x in args.bug_id.split(",") if x.strip()]
        except ValueError:
            print(f"❌ Bug ID 格式错误: {args.bug_id}（应为逗号分隔的数字）", file=sys.stderr)
            sys.exit(1)
        if not bug_ids:
            print("❌ 未提供有效 Bug ID", file=sys.stderr)
            sys.exit(1)
        bugs = []
        for bid in bug_ids:
            try:
                print(f"  从禅道拉取 Bug #{bid}...")
                bug = zentao.fetch_bug_as_bugitem(bid)
                bugs.append(bug)
                print(f"  ✅ #{bid}: {bug.title}")
            except Exception as e:
                print(f"  ❌ Bug #{bid} 拉取失败: {e}", file=sys.stderr)
                sys.exit(1)
    else:
        # 验证 Excel
        if not Path(args.excel).is_file():
            print(f"❌ Excel 文件不存在: {args.excel}", file=sys.stderr)
            sys.exit(1)
        bugs = load_bugs_from_excel(args.excel)
        if not bugs:
            print(f"❌ Excel 中无有效 bug 行", file=sys.stderr)
            sys.exit(1)

    print(f"加载 {len(bugs)} 个 bug")

    # 创建适配器
    try:
        adapter = ZCodeAdapter(provider_hint=args.provider)
    except Exception as e:
        print(f"❌ ZCode 适配器初始化失败: {e}", file=sys.stderr)
        sys.exit(1)

    if not adapter.health_check():
        print(f"❌ ZCode CLI 不可用，请检查安装和凭证配置", file=sys.stderr)
        sys.exit(1)

    # 运行
    orchestrator = BugfixOrchestrator(
        adapter=adapter,
        server=args.server,
        code_root=args.code_root,
        device_serial=args.device_serial,
        timeout_per_session=args.timeout,
        dry_run=args.dry_run,
        legacy_3session=args.legacy_3session,
        resume=args.resume,
    )

    bugs = orchestrator.run_batch(bugs)

    # 回写 Excel（仅 --excel 模式；--bug-id 模式无 Excel 文件）
    if not args.dry_run and args.excel:
        try:
            write_results_to_excel(args.excel, bugs)
            print(f"\n结果已回写到 {args.excel}")
        except Exception as e:
            print(f"\n⚠️ 回写 Excel 失败: {e}", file=sys.stderr)


if __name__ == "__main__":
    main()
